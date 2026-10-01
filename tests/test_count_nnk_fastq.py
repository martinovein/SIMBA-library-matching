import gzip
import sys
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "scripts"))

import count_nnk_fastq as c


def make_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "NNK_library"
    ws.append(["parent_id", "oligo_id", "wt_peptide_aa", "wt_coding_dna", "nnk_coding_pattern", "randomized_position_aa", "fixed_5prime_sequence", "fixed_3prime_sequence"])
    ws.append(["P1", "P1-WT", "MK", "ATGAAA", "ATGAAA", "WT", "caatgacgtcTCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGGAgctagcaccg"])
    ws.append(["P1", "P1-NNK1", "MK", "ATGAAA", "NNKAAA", 1, "caatgacgtcTCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGGAgctagcaccg"])
    ws.append(["P1", "P1-NNK2", "MK", "ATGAAA", "ATGNNK", 2, "caatgacgtcTCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGGAgctagcaccg"])
    ws = wb.create_sheet("Fixed_controls")
    ws.append(["control_id", "control_type", "peptide_aa", "coding_dna", "fixed_5prime_sequence", "fixed_3prime_sequence", "expected_function", "note"])
    ws.append(["ctrl1", "negative", "GP", "GGTCCA", "caatgacgtcTCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGGAgctagcaccg", "none", "extra metadata ignored by parser"])
    ws = wb.create_sheet("Samples")
    ws.append(["fastq_filename", "sample_id", "replicate", "bait", "time_hours", "is_common_input"])
    ws.append(["s1.fastqsanger.gz", "s1", "rep1", "B", 0, False])
    wb.save(path)


def test_nnk_expansion_exactly_32_and_k_restriction(tmp_path):
    wb = tmp_path / "lib.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    (data / "s1.fastqsanger.gz").write_bytes(b"")
    make_workbook(wb)
    bundle = c.validate_and_expand(wb, data, None, None)
    assert not bundle.fatal_messages
    pos1 = [v.observed_codon for v in bundle.dna_variants.values() if v.randomized_position_aa == 1]
    assert len(pos1) == 32
    assert sorted(pos1) == sorted(c.NNK_CODONS)
    assert {codon[2] for codon in pos1} == {"G", "T"}


def test_translation_and_randomized_position_validation(tmp_path):
    assert c.translate_dna("ATGTAA") == "M*"
    wb = tmp_path / "bad.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    (data / "s1.fastqsanger.gz").write_bytes(b"")
    make_workbook(wb)
    # Corrupt randomized position.
    from openpyxl import load_workbook
    owb = load_workbook(wb)
    owb["NNK_library"]["F3"] = 2
    owb.save(wb)
    bundle = c.validate_and_expand(wb, data, None, None)
    assert any("randomized_position_aa" in msg for msg in bundle.fatal_messages)


def test_duplicate_reference_aliases(tmp_path):
    wb = tmp_path / "lib.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    (data / "s1.fastqsanger.gz").write_bytes(b"")
    make_workbook(wb)
    from openpyxl import load_workbook
    owb = load_workbook(wb)
    owb["NNK_library"].append(["P1", "P1-NNK1-copy", "MK", "ATGAAA", "NNKAAA", 1, "caatgacgtcTCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGGAgctagcaccg"])
    owb.save(wb)
    bundle = c.validate_and_expand(wb, data, None, None)
    wt = bundle.seq_to_variant["ATGAAA"]
    assert len(bundle.dna_variants[wt].aliases) == 2


def test_forward_reverse_low_quality_missing_and_length_mapping():
    a5, a3 = "ACGTACGTACGTACG", "TGCATGCATGCATGC"
    seq = a5 + "ATGAAA" + a3
    qual = "I" * len(seq)
    params = dict(min_base_quality=20, min_mean_quality=None, max_n=0, expected_lengths={6}, seq_to_variant={"ATGAAA": "v1"}, seq_to_control={})
    ext = c.extract_read(seq, qual, a5, a3)
    assert ext.orientation == "forward"
    assert c.classify_extracted(ext, **params) == ("mapped_nnk", "ATGAAA", "v1")
    rc = c.reverse_complement(seq)
    ext = c.extract_read(rc, qual[::-1], a5, a3)
    assert ext.orientation == "reverse_complement"
    low = c.extract_read(seq, "!" * len(seq), a5, a3)
    assert c.classify_extracted(low, **params)[0] == "low_quality"
    assert c.classify_extracted(c.extract_read("TTTT", "IIII", a5, a3), **params)[0] == "missing_flank"
    bad_len = c.extract_read(a5 + "ATGA" + a3, "I" * (len(a5) + 4 + len(a3)), a5, a3)
    assert c.classify_extracted(bad_len, **params)[0] == "invalid_length_frame"


def test_unexpected_multiple_mutated_codon_and_control_assignment():
    a5, a3 = "ACGTACGTACGTACG", "TGCATGCATGCATGC"
    params = dict(min_base_quality=20, min_mean_quality=None, max_n=0, expected_lengths={6}, seq_to_variant={"ATGAAA": "v1"}, seq_to_control={"GGTCCA": "ctrl"})
    ctrl = c.extract_read(a5 + "GGTCCA" + a3, "I" * (len(a5) + 6 + len(a3)), a5, a3)
    assert c.classify_extracted(ctrl, **params) == ("mapped_control", "GGTCCA", "ctrl")
    unexp = c.extract_read(a5 + "TTTCCC" + a3, "I" * (len(a5) + 6 + len(a3)), a5, a3)
    assert c.classify_extracted(unexp, **params)[0] == "unexpected"
    assert c.altered_codons("TTTCCC", "ATGAAA") == 2


def test_dna_to_protein_aggregation_and_enrich_format(tmp_path):
    v1 = c.DnaVariant("P:c.1A>G", "P", "GTGAAA", "VK", "c.1A>G", "p.Met1Val", "missense", 1, "ATG", "GTG", "M", "V")
    wt = c.DnaVariant("P:_wt", "P", "ATGAAA", "MK", "_wt", "_wt", "WT", 1, "ATG", "ATG", "M", "M")
    stop = c.DnaVariant("P:c.4A>T", "P", "ATGTAA", "M*", "c.4A>T", "p.Lys2Ter", "stop", 2, "AAA", "TAA", "K", "*")
    assert c.enrich_variant_name(v1) == "c.1A>G"
    assert c.enrich_variant_name(wt) == "_wt"
    assert c.protein_variant_id(wt) == "P_WT"
    assert c.protein_variant_id(v1) == "P_M1V"
    assert c.protein_variant_id(stop) == "P_K2*"
    assert c.protein_hgvs("MK", "MK") == "_wt"


def test_missing_unexpected_fastq_detection(tmp_path):
    wb = tmp_path / "lib.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    make_workbook(wb)
    (data / "extra.fastqsanger.gz").write_bytes(b"")
    bundle = c.validate_and_expand(wb, data, None, None)
    assert any("Missing FASTQ" in msg for msg in bundle.fatal_messages)
    assert any("Unexpected FASTQ" in msg for msg in bundle.warnings)


def test_synthetic_fastq_integration(tmp_path):
    wb = tmp_path / "lib.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    make_workbook(wb)
    a5, a3 = "TATGCAGCCTCTTCATCTGGC", "GGTGGAGGATCCGGAGG"
    forward = a5 + "ATGAAA" + a3
    control = a5 + "GGTCCA" + a3
    reverse = c.reverse_complement(forward)
    low = a5 + "ATGAAA" + a3
    unexpected = a5 + "TTTCCC" + a3
    bad = a5 + "ATGA" + a3
    records = [
        ("@f", forward, "I" * len(forward)),
        ("@r", reverse, "I" * len(reverse)),
        ("@l", low, "!" * len(low)),
        ("@c", control, "I" * len(control)),
        ("@u", unexpected, "I" * len(unexpected)),
        ("@b", bad, "I" * len(bad)),
    ]
    with gzip.open(data / "s1.fastqsanger.gz", "wt") as fh:
        for h, s, q in records:
            fh.write(f"{h}\n{s}\n+\n{q}\n")
    bundle = c.validate_and_expand(wb, data, None, None)
    params = {
        "anchor5": bundle.anchor5,
        "anchor3": bundle.anchor3,
        "expected_lengths": bundle.expected_lengths,
        "min_base_quality": 20,
        "min_mean_quality": None,
        "max_n": 0,
        "max_reads_per_file": None,
        "seq_to_variant": bundle.seq_to_variant,
        "seq_to_control": bundle.seq_to_control,
    }
    res = c.process_sample_worker((bundle.samples[0].__dict__, str(data), params))
    assert res["total_reads"] == 6
    assert sum(res["category_counts"].get(cat, 0) for cat in c.READ_CATEGORIES) == 6
    assert res["category_counts"]["mapped_nnk"] == 2
    assert res["category_counts"]["mapped_control"] == 1
    assert res["category_counts"]["low_quality"] == 1
    assert res["category_counts"]["unexpected"] == 1
    assert res["category_counts"]["invalid_length_frame"] == 1


def test_common_input_nonnumeric_time_and_checkpoint_fingerprint(tmp_path):
    wb = tmp_path / "lib.xlsx"
    data = tmp_path / "data"
    data.mkdir()
    make_workbook(wb)
    (data / "s1.fastqsanger.gz").write_bytes(b"")
    from openpyxl import load_workbook
    owb = load_workbook(wb)
    ws = owb["Samples"]
    ws.append(["input.fastqsanger.gz", "input", "input", "input", "input", True])
    owb.save(wb)
    (data / "input.fastqsanger.gz").write_bytes(b"")
    bundle = c.validate_and_expand(wb, data, None, None)
    assert not bundle.fatal_messages
    common = [s for s in bundle.samples if s.is_common_input][0]
    assert common.time_hours is None
    assert any("nonnumeric time_hours" in w for w in bundle.warnings)
    params = {
        "workbook_sha256": c.sha256_file(wb),
        "anchor5": bundle.anchor5,
        "anchor3": bundle.anchor3,
        "min_base_quality": 20,
        "min_mean_quality": None,
        "max_n": 0,
        "max_reads_per_file": 10,
    }
    fp = c.checkpoint_fingerprint(common, data, params)
    assert fp["workbook_sha256"] == c.sha256_file(wb)
    assert fp["pipeline_version"] == c.PIPELINE_VERSION
    assert fp["reference_expansion_version"] == c.REFERENCE_EXPANSION_VERSION
